from flask import Flask, render_template, request, session
import os
import uuid
import time
import speech_recognition as sr
from deep_translator import GoogleTranslator
from gtts import gTTS
from pydub import AudioSegment
import PyPDF2
from docx import Document
import openpyxl
from PIL import Image
import pytesseract

# ---------------- PATH SETTINGS ----------------
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
AudioSegment.converter = r"C:\ffmpeg\bin\ffmpeg.exe"
AudioSegment.ffprobe = r"C:\ffmpeg\bin\ffprobe.exe"

app = Flask(__name__)
app.secret_key = "voice_project_123"

UPLOAD_FOLDER = "uploads"
OUTPUT_FOLDER = "static"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)


# ---------------- TEXT SPLITTER ----------------
def split_text(text, chunk_size=2000):
    return [text[i:i + chunk_size] for i in range(0, len(text), chunk_size)]


# ---------------- TRANSLATION ----------------
def translate_text(text, target_lang):
    chunks = split_text(text)
    translated_chunks = []

    for chunk in chunks:
        translated = GoogleTranslator(
            source="auto",
            target=target_lang
        ).translate(chunk)

        translated_chunks.append(translated)
        time.sleep(0.3)

    return " ".join(translated_chunks)


# ---------------- TEXT TO AUDIO ----------------
def text_to_audio(text, lang):
    unique_id = uuid.uuid4().hex
    audio_parts = []

    chunks = split_text(text)

    for i, chunk in enumerate(chunks):
        part_path = os.path.join(OUTPUT_FOLDER, f"part_{unique_id}_{i}.mp3")
        gTTS(chunk, lang=lang).save(part_path)
        audio_parts.append(part_path)

    final_audio = AudioSegment.empty()

    for part in audio_parts:
        final_audio += AudioSegment.from_mp3(part)

    output_file = f"output_{unique_id}.mp3"
    final_audio.export(os.path.join(OUTPUT_FOLDER, output_file), format="mp3")

    return output_file


# ---------------- HOME ----------------
@app.route("/")
def home():
    return render_template("index.html")


@app.route("/voice")
def voice_page():
    return render_template("voice.html")


@app.route("/pdf")
def pdf_page():
    return render_template("pdf.html")


# ---------------- MULTILINGUAL VOICE CONVERT ----------------
@app.route("/voice_convert", methods=["POST"])
def voice_convert():

    if "audio" not in request.files:
        return "No audio recorded!"

    audio_file = request.files["audio"]
    target_lang = request.form.get("language")

    uid = uuid.uuid4().hex
    webm_path = os.path.join(UPLOAD_FOLDER, f"{uid}.webm")
    wav_path = os.path.join(UPLOAD_FOLDER, f"{uid}.wav")

    audio_file.save(webm_path)

    sound = AudioSegment.from_file(webm_path, format="webm")
    sound.export(wav_path, format="wav")

    recognizer = sr.Recognizer()

    with sr.AudioFile(wav_path) as source:
        audio_data = recognizer.record(source)

    # 🌍 MULTI LANGUAGE SPEECH DETECTION
    languages = [
        "en-US",
        "hi-IN",
        "fr-FR",
        "de-DE",
        "ja-JP",
        "ru-RU",
        "es-ES",
        "zh-CN"
    ]

    text = None

    for lang in languages:
        try:
            text = recognizer.recognize_google(audio_data, language=lang)
            print("Detected Language:", lang)
            break
        except:
            continue

    if text is None:
        return "Speech not recognized."

    session["original_text"] = text
    session["source_type"] = "voice"

    translated_text = translate_text(text, target_lang)
    audio_output = text_to_audio(translated_text, target_lang)

    return render_template(
        "voice.html",
        original_text=text,
        translated_text=translated_text,
        audio_file=audio_output
    )


# ---------------- DOCUMENT CONVERT ----------------
@app.route("/doc_convert", methods=["POST"])
def doc_convert():

    file = request.files["doc_file"]
    language = request.form["language"]
    feature = request.form["feature"]

    file_ext = file.filename.split(".")[-1].lower()
    file_path = os.path.join(UPLOAD_FOLDER, file.filename)
    file.save(file_path)

    extracted_text = ""

    if file_ext == "pdf":
        reader = PyPDF2.PdfReader(open(file_path, "rb"))
        for page in reader.pages:
            extracted_text += page.extract_text() or ""

    elif file_ext == "docx":
        doc = Document(file_path)
        for para in doc.paragraphs:
            extracted_text += para.text + "\n"

    elif file_ext == "xlsx":
        wb = openpyxl.load_workbook(file_path)
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell:
                        extracted_text += str(cell) + " "

    elif file_ext in ["png", "jpg", "jpeg"]:
        img = Image.open(file_path)
        extracted_text = pytesseract.image_to_string(img)

    if not extracted_text.strip():
        return "No readable text found."

    session["original_text"] = extracted_text
    session["source_type"] = "document"

    translated_text = None
    audio_file = None

    if feature == "translate":
        translated_text = translate_text(extracted_text, language)

    elif feature == "audio":
        audio_file = text_to_audio(extracted_text, language)

    elif feature == "all":
        translated_text = translate_text(extracted_text, language)
        audio_file = text_to_audio(translated_text, language)

    return render_template(
        "pdf.html",
        original_text=extracted_text,
        translated_text=translated_text,
        audio_file=audio_file
    )


# ---------------- RECONVERT ----------------
@app.route("/reconvert", methods=["POST"])
def reconvert():

    language = request.form["language"]
    feature = request.form["feature"]

    original_text = session.get("original_text")
    source_type = session.get("source_type")

    translated_text = None
    audio_file = None

    if feature == "translate":
        translated_text = translate_text(original_text, language)

    elif feature == "audio":
        audio_file = text_to_audio(original_text, language)

    elif feature == "all":
        translated_text = translate_text(original_text, language)
        audio_file = text_to_audio(translated_text, language)

    page = "voice.html" if source_type == "voice" else "pdf.html"

    return render_template(
        page,
        original_text=original_text,
        translated_text=translated_text,
        audio_file=audio_file
    )


if __name__ == "__main__":
    app.run(debug=True)